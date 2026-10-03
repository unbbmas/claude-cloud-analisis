import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EMPRESA = "Nichiwa Sangyo Co"
TICKER = "2055"
FECHA = "3/10/2026"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)

# Abreviaturas de fuentes primarias (se citan en cada dato)
YUHO = {
    2021: "Yuho FY3/2021 (F_2020.pdf, presentado jun-2021)",
    2022: "Yuho FY3/2022 (F_2021.pdf, presentado jun-2022)",
    2023: "Yuho FY3/2023 (F_2022.pdf, presentado jun-2023)",
    2024: "Yuho FY3/2024 (F_2023.pdf, presentado jun-2024)",
    2025: "Yuho FY3/2025 (F_2024.pdf, presentado jun-2025)",
    2026: "Yuho FY3/2026 (F_2025.pdf, 25-jun-2026)",
}
Q1 = "Kessan Tanshin 1T FY3/2027 (Q1_2026.pdf, 12-ago-2026)"


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _runs(par, text, size=None, color=None, italic=False):
    """Admite **negrita** dentro del texto."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for p in parts:
        if not p:
            continue
        bold = p.startswith("**") and p.endswith("**")
        r = par.add_run(p[2:-2] if bold else p)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return par


class Doc:
    def __init__(self, num, titulo):
        self.num = num
        self.titulo = titulo
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
        for s in self.d.sections:
            s.left_margin = s.right_margin = Cm(2.0)
            s.top_margin = s.bottom_margin = Cm(1.8)
        t = self.d.add_heading(level=0)
        _runs(t, f"{titulo} - {EMPRESA} ({TICKER}) - {FECHA}")
        p = self.d.add_paragraph()
        _runs(p, f"Análisis de Inversión - {EMPRESA} ({TICKER}) - {FECHA} · Sección {num} de 12", size=9, color=NAVY, italic=True)
        p = self.d.add_paragraph()
        _runs(p, "Moneda: **JPY** (millones de yenes salvo indicación). Ejercicio fiscal cerrado a 31 de marzo (FY3/2026 = abr-2025 a mar-2026). "
                 "Precio de referencia: **346 JPY** (cierre del último día hábil, 2-oct-2026; verificado en rango 345-346 JPY a 25-sep-2026). "
                 "Normas contables: J-GAAP. Leyenda: 🔍 DATO VERIFICADO · 💡 ESTIMACIÓN · ⚠️ ESPECULACIÓN · 🟢 Positivo / 🟡 Neutral / 🔴 Negativo.",
              size=9)

    def h1(self, text):
        self.d.add_heading(text, level=1)

    def h2(self, text):
        self.d.add_heading(text, level=2)

    def h3(self, text):
        self.d.add_heading(text, level=3)

    def p(self, text, size=None, italic=False):
        par = self.d.add_paragraph()
        _runs(par, text, size=size, italic=italic)
        return par

    def bullets(self, items):
        for it in items:
            par = self.d.add_paragraph(style="List Bullet")
            _runs(par, it)

    def numbered(self, items):
        for it in items:
            par = self.d.add_paragraph(style="List Number")
            _runs(par, it)

    def src(self, text):
        par = self.d.add_paragraph()
        _runs(par, f"[Fuente: {text}]", size=8.5, color=RGBColor(0x55, 0x55, 0x55), italic=True)

    def calc(self, text):
        par = self.d.add_paragraph()
        _runs(par, "💡 ESTIMACIÓN / CÁLCULO PROPIO: " + text, size=9.5, color=RGBColor(0x1F, 0x4E, 0x79))

    def box(self, text, fill="EAF1FB"):
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.cell(0, 0)
        _shade(c, fill)
        c.paragraphs[0].text = ""
        for i, line in enumerate(text.split("\n")):
            par = c.paragraphs[0] if i == 0 else c.add_paragraph()
            _runs(par, line)
        self.d.add_paragraph()

    def table(self, header, rows, widths=None, size=8.5, first_bold=True):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            _shade(c, "1F3A5F")
            c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(str(h))
            r.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for row in rows:
            cells = t.add_row().cells
            section_row = len(row) == 1
            if section_row:
                m = cells[0].merge(cells[-1])
                _shade(m, "D9E2F3")
                m.paragraphs[0].text = ""
                _runs(m.paragraphs[0], f"**{row[0]}**", size=size)
                continue
            for i, v in enumerate(row):
                cells[i].paragraphs[0].text = ""
                txt = str(v)
                if i == 0 and first_bold and not txt.startswith("**"):
                    txt = f"**{txt}**"
                _runs(cells[i].paragraphs[0], txt, size=size)
                if i > 0:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph()
        return t

    def save(self, folder, slug):
        path = f"{folder}/{self.num:02d}_{slug}_Nichiwa_Sangyo_2055.docx"
        self.d.save(path)
        return path


def n(x, dec=0):
    """Formato numérico español: miles con punto, decimales con coma."""
    if x is None:
        return "n.d."
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, dec=1):
    return n(x, dec) + "%"
